"""Seed the demo user + the Verdure Materials workspace, with its brand corpus
pre-ingested (DECISIONS #15 — visitors hit the RAG "wow" instantly).

Run from the server/ directory:  python -m seed.seed_demo
(or `npm run seed` from the repo root). Idempotent: re-running won't duplicate
the user, workspace, or assets. Corpus ingestion needs OPENAI_API_KEY + a
running Chroma; if the key is missing the workspace is created empty and you can
re-run after adding the key.
"""
import asyncio
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from bson import ObjectId
from motor.motor_asyncio import AsyncIOMotorClient

from app.config import settings
from app.rag.loaders import load_sections
from app.rag.vector_store import get_vector_store
from app.services.ingestion_service import build_chunks
from app.services.storage_service import build_image_key, build_storage_key, build_template_key, get_storage
from app.services.template_service import discover_placeholders
from app.utils.security import get_password_hash

CORPUS_DIR = Path(__file__).resolve().parent / "verdure_corpus"

# Verdure brand palette (RGB tuples for Pillow; hex for Office docs).
_FOREST = (0x1C, 0x3A, 0x2E)
_CREAM = (0xF4, 0xF1, 0xE9)
_OCHRE = (0xC9, 0x8A, 0x3D)
_FOREST_HEX = "1C3A2E"
_OCHRE_HEX = "C98A3D"

# (filename, content_type, asset_kind). The voice guide is ingested as a PDF to
# exercise the pdfplumber path; the rest are markdown.
CORPUS = [
    ("voice-and-tone.pdf", "application/pdf", "voice"),
    ("content-pillars.md", "text/markdown", "pillars"),
    ("past-posts.md", "text/markdown", "posts"),
    ("personas.md", "text/markdown", "personas"),
    ("competitors.md", "text/markdown", "competitors"),
]

VERDURE_BRAND_PROFILE = {
    "colors": [
        {"name": "Forest", "hex": "#1C3A2E", "role": "primary"},
        {"name": "Cream", "hex": "#F4F1E9", "role": "neutral"},
        {"name": "Ochre", "hex": "#C98A3D", "role": "accent"},
    ],
    "accent_hex": "#1C3A2E",
    "fonts": {"heading": "DM Serif Display", "body": "Spectral", "mono": None},
    "logo": None,
    "voice_summary": (
        "Verdure is the grounded, data-first voice of low-carbon building materials — "
        "plain-spoken, evidence-led, and respectful of the professionals who specify and build."
    ),
}


async def _ensure_demo_user(db) -> ObjectId:
    demo = await db.users.find_one({"is_demo": True})
    if demo:
        print(f"✓ demo user exists ({demo['email']})")
        return demo["_id"]
    now = datetime.now(timezone.utc)
    res = await db.users.insert_one(
        {
            "email": settings.demo_email,
            "hashed_password": get_password_hash(settings.demo_password),
            "full_name": settings.demo_full_name,
            "is_active": True,
            "is_admin": False,
            "is_demo": True,
            "created_at": now,
            "updated_at": now,
        }
    )
    print(f"+ created demo user ({settings.demo_email})")
    return res.inserted_id


async def _ensure_workspace(db, owner_id: ObjectId) -> dict:
    existing = await db.workspaces.find_one({"owner_id": owner_id, "slug": "verdure-materials"})
    if existing:
        print("✓ Verdure Materials workspace exists")
        return existing
    now = datetime.now(timezone.utc)
    workspace_oid = ObjectId()
    doc = {
        "_id": workspace_oid,
        "owner_id": owner_id,
        "name": "Verdure Materials",
        "slug": "verdure-materials",
        "brand_profile": VERDURE_BRAND_PROFILE,
        "vector_namespace": str(workspace_oid),
        "embedding_model": settings.embedding_model,
        "embedding_dim": settings.embedding_dim,
        "asset_count": 0,
        "ready_chunk_count": 0,
        "created_at": now,
        "updated_at": now,
    }
    await db.workspaces.insert_one(doc)
    print("+ created Verdure Materials workspace")
    return doc


async def _ingest_corpus(db, workspace: dict) -> None:
    if await db.assets.count_documents({"workspace_id": workspace["_id"]}) > 0:
        print("✓ corpus already ingested")
        return
    if not settings.openai_api_key.get_secret_value():
        print("⚠ OPENAI_API_KEY not set — workspace created empty. Add the key and re-run to ingest.")
        return

    namespace = workspace["vector_namespace"]
    store = get_vector_store()
    storage = get_storage()
    total_chunks = 0
    asset_count = 0
    now = datetime.now(timezone.utc)

    for filename, content_type, kind in CORPUS:
        path = CORPUS_DIR / filename
        if not path.exists():
            print(f"⚠ missing corpus file {filename}; skipping")
            continue
        data = path.read_bytes()
        asset_oid = ObjectId()
        storage_key = build_storage_key(namespace, str(asset_oid), filename)
        await storage.save(storage_key, data)

        sections = load_sections(storage.local_path(storage_key), content_type)
        chunks = build_chunks(
            workspace_namespace=namespace,
            asset_id=str(asset_oid),
            asset_kind=kind,
            filename=filename,
            sections=sections,
        )
        count = await store.upsert_chunks(namespace, chunks)
        await db.assets.insert_one(
            {
                "_id": asset_oid,
                "workspace_id": workspace["_id"],
                "filename": filename,
                "content_type": content_type,
                "asset_kind": kind,
                "storage_key": storage_key,
                "size_bytes": len(data),
                "ingestion_status": "ready",
                "chunk_count": count,
                "error": None,
                "ingested_at": now,
                "created_at": now,
                "updated_at": now,
            }
        )
        total_chunks += count
        asset_count += 1
        print(f"  + {filename} → {count} chunks ({kind})")

    await db.workspaces.update_one(
        {"_id": workspace["_id"]},
        {"$set": {"asset_count": asset_count, "ready_chunk_count": total_chunks, "updated_at": now}},
    )
    print(f"+ ingested Verdure corpus: {total_chunks} chunks across {asset_count} assets")


# ---------------------------------------------------------------------------
# Demo content generators (deterministic — no model calls, no external assets)
# ---------------------------------------------------------------------------
def _png_bytes(img) -> bytes:  # type: ignore[no-untyped-def]
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _font(size: int):  # type: ignore[no-untyped-def]
    from PIL import ImageFont

    return ImageFont.load_default(size=size)


def _draw_wrapped(draw, text, font, fill, *, x, y, max_width, line_h):  # type: ignore[no-untyped-def]
    words = text.split()
    line = ""
    for word in words:
        trial = f"{line} {word}".strip()
        if draw.textlength(trial, font=font) <= max_width or not line:
            line = trial
        else:
            draw.text((x, y), line, font=font, fill=fill)
            y += line_h
            line = word
    if line:
        draw.text((x, y), line, font=font, fill=fill)


def _generate_logo() -> bytes:
    """A simple Verdure wordmark — forest field, cream text, ochre rule."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (480, 480), _FOREST)
    d = ImageDraw.Draw(img)
    d.text((60, 200), "Verdure", font=_font(72), fill=_CREAM)
    d.rectangle([60, 300, 300, 308], fill=_OCHRE)
    d.text((62, 330), "MATERIALS", font=_font(34), fill=_OCHRE)
    return _png_bytes(img)


def _generate_card(eyebrow: str, headline: str) -> bytes:
    """A 1080×1350 Verdure-styled card standing in for a past brand post."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1080, 1350), _FOREST)
    d = ImageDraw.Draw(img)
    d.text((90, 140), eyebrow.upper(), font=_font(40), fill=_OCHRE)
    _draw_wrapped(d, headline, _font(92), _CREAM, x=90, y=320, max_width=900, line_h=112)
    d.rectangle([90, 1180, 230, 1188], fill=_OCHRE)
    d.text((90, 1210), "verdure-materials.example", font=_font(30), fill=_CREAM)
    return _png_bytes(img)


def _generate_brand_image(label: str) -> bytes:
    """A deterministic branded swatch for the demo brand-image library."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1080, 1350), _FOREST)
    d = ImageDraw.Draw(img)
    d.rectangle([60, 60, 1020, 1290], outline=_OCHRE, width=8)
    d.text((90, 1190), label.upper(), font=_font(44), fill=_CREAM)
    return _png_bytes(img)


def _generate_docx_template() -> bytes:
    """A Verdure-branded one-pager with {{placeholders}} to fill."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Spectral"
    normal.font.size = Pt(11)

    def add(text, *, size, color_hex, bold=False, upper=False):
        para = doc.add_paragraph()
        run = para.add_run(text.upper() if upper else text)
        run.font.name = "DM Serif Display" if bold else "Spectral"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color_hex)
        return para

    add("VERDURE MATERIALS", size=10, color_hex=_OCHRE_HEX, upper=True)
    add("{{project_title}}", size=24, color_hex=_FOREST_HEX, bold=True)
    add("Prepared for {{client_name}} — {{date}}", size=11, color_hex="6C737A")
    add("{{summary}}", size=12, color_hex="1A1D21")
    add("Next steps", size=15, color_hex=_FOREST_HEX, bold=True)
    add("{{next_steps}}", size=11, color_hex="1A1D21")
    add("Prepared by {{prepared_by}}", size=9, color_hex="6C737A")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _generate_pptx_template() -> bytes:
    """A Verdure-branded 16:9 deck (title + content slide) with {{placeholders}}."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor as PptxRGB
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_line(tf, text, *, size, color_hex, bold=False, upper=False, first=False):
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        run = para.add_run()
        run.text = text.upper() if upper else text
        run.font.name = "DM Serif Display" if bold else "Spectral"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = PptxRGB.from_string(color_hex)

    title = prs.slides.add_slide(blank).shapes.add_textbox(
        Inches(0.9), Inches(2.4), Inches(11.5), Inches(3)
    ).text_frame
    title.word_wrap = True
    add_line(title, "VERDURE MATERIALS", size=18, color_hex=_OCHRE_HEX, upper=True, first=True)
    add_line(title, "{{title}}", size=44, color_hex=_FOREST_HEX, bold=True)
    add_line(title, "{{subtitle}}", size=20, color_hex="1A1D21")

    content = prs.slides.add_slide(blank).shapes.add_textbox(
        Inches(0.9), Inches(0.8), Inches(11.5), Inches(5.9)
    ).text_frame
    content.word_wrap = True
    add_line(content, "{{section_title}}", size=30, color_hex=_FOREST_HEX, bold=True, first=True)
    add_line(content, "{{point_1}}", size=20, color_hex="1A1D21")
    add_line(content, "{{point_2}}", size=20, color_hex="1A1D21")
    add_line(content, "{{point_3}}", size=20, color_hex="1A1D21")

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Demo enrichment (idempotent; Verdure-only). Demo mutations are blocked, so we
# seed the Phase 4–5 surfaces here or they'd only ever show empty states.
# ---------------------------------------------------------------------------
async def _ensure_logo(db, workspace: dict) -> None:
    ws_id = str(workspace["_id"])
    fresh = await db.workspaces.find_one({"_id": workspace["_id"]})
    if (fresh.get("brand_profile") or {}).get("logo"):
        print("✓ brand logo exists")
        return
    logo_id = str(ObjectId())
    name = "verdure-logo.png"
    key = build_image_key(ws_id, logo_id, name)
    await get_storage().save(key, _generate_logo())
    logo = {"storage_key": key, "url": f"/api/workspaces/{ws_id}/images/file/{logo_id}/{name}"}
    await db.workspaces.update_one(
        {"_id": workspace["_id"]},
        {"$set": {"brand_profile.logo": logo, "updated_at": datetime.now(timezone.utc)}},
    )
    print("+ seeded brand logo")


async def _ensure_exemplars(db, workspace: dict) -> None:
    fresh = await db.workspaces.find_one({"_id": workspace["_id"]})
    if (fresh.get("brand_profile") or {}).get("exemplars"):
        print("✓ example cards exist")
        return
    ws_id = str(workspace["_id"])
    cards = [
        (
            "Material notes",
            "The carbon math is changing",
            "Punchy declarative headline (<=8 words) over a deep forest field; ochre eyebrow kicker; one tight idea per card; calm, evidence-led tone.",
        ),
        (
            "Field guide",
            "Specify hempcrete early",
            "Imperative how-to headline; generous negative space; cream type on forest; ends with a quiet, practical takeaway, never a hard sell.",
        ),
    ]
    exemplars = []
    for eyebrow, headline, note in cards:
        ex_id = str(ObjectId())
        name = f"exemplar-{ex_id}.png"
        key = build_image_key(ws_id, ex_id, name)
        await get_storage().save(key, _generate_card(eyebrow, headline))
        exemplars.append(
            {
                "id": ex_id,
                "storage_key": key,
                "url": f"/api/workspaces/{ws_id}/images/file/{ex_id}/{name}",
                "style_note": note,
            }
        )
    await db.workspaces.update_one(
        {"_id": workspace["_id"]},
        {"$set": {"brand_profile.exemplars": exemplars, "updated_at": datetime.now(timezone.utc)}},
    )
    print(f"+ seeded {len(exemplars)} example cards")


async def _ensure_brand_images(db, workspace: dict) -> None:
    fresh = await db.workspaces.find_one({"_id": workspace["_id"]})
    if (fresh.get("brand_profile") or {}).get("brand_images"):
        print("✓ brand images exist")
        return
    ws_id = str(workspace["_id"])
    images = []
    for label in ("Forest texture", "Ochre accent"):
        img_id = str(ObjectId())
        name = f"{label.lower().replace(' ', '-')}.png"
        key = build_image_key(ws_id, img_id, name)
        await get_storage().save(key, _generate_brand_image(label))
        images.append(
            {
                "id": img_id,
                "storage_key": key,
                "url": f"/api/workspaces/{ws_id}/images/file/{img_id}/{name}",
                "name": name,
            }
        )
    await db.workspaces.update_one(
        {"_id": workspace["_id"]},
        {"$set": {"brand_profile.brand_images": images, "updated_at": datetime.now(timezone.utc)}},
    )
    print(f"+ seeded {len(images)} brand images")


async def _ensure_templates(db, workspace: dict) -> None:
    if await db.templates.count_documents({"workspace_id": workspace["_id"]}) > 0:
        print("✓ templates exist")
        return
    ws_id = str(workspace["_id"])
    now = datetime.now(timezone.utc)
    files = [
        ("Verdure one-pager.docx", "docx", _generate_docx_template()),
        ("Verdure deck.pptx", "pptx", _generate_pptx_template()),
    ]
    for name, fmt, data in files:
        template_oid = ObjectId()
        key = build_template_key(ws_id, str(template_oid), name)
        await get_storage().save(key, data)
        placeholders = discover_placeholders(data, fmt)
        await db.templates.insert_one(
            {
                "_id": template_oid,
                "workspace_id": workspace["_id"],
                "owner_id": workspace["owner_id"],
                "name": name.replace(" ", "_"),
                "format": fmt,
                "storage_key": key,
                "placeholders": placeholders,
                "created_at": now,
            }
        )
        print(f"  + {name} → {len(placeholders)} fields ({fmt})")
    print("+ seeded document templates")


async def main() -> None:
    client = AsyncIOMotorClient(settings.mongodb_url.get_secret_value())
    db = client[settings.database_name]
    await db.users.create_index("email", unique=True)

    owner_id = await _ensure_demo_user(db)
    workspace = await _ensure_workspace(db, owner_id)
    await _ingest_corpus(db, workspace)
    await _ensure_logo(db, workspace)
    await _ensure_exemplars(db, workspace)
    await _ensure_brand_images(db, workspace)
    await _ensure_templates(db, workspace)

    client.close()
    print("Done.")


if __name__ == "__main__":
    asyncio.run(main())
