"""On-brand document templates (Phase 5). Discover the `{{placeholder}}` tags in
an uploaded .docx/.pptx, fill them with values (the user's own text or an AI
draft), and render — swapping ONLY text, never the template's styling. docx via
docxtpl (Jinja2-for-Word); pptx via paragraph-level run replacement (PowerPoint
splits a tag across runs, so we substitute on the joined paragraph text)."""
from __future__ import annotations

import re
from io import BytesIO

from bson import ObjectId
from docxtpl import DocxTemplate
from jinja2.sandbox import SandboxedEnvironment
from motor.motor_asyncio import AsyncIOMotorDatabase
from pptx import Presentation

from app.models.common import serialize_doc, utcnow
from app.models.template import TemplateResponse
from app.utils.archives import assert_safe_ooxml

MAX_PLACEHOLDERS = 40
MAX_VALUE_LEN = 2000
_TAG = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


class TemplateHasLogic(ValueError):
    """An uploaded template uses Jinja statements/comments ({% %} / {# #}); only
    {{placeholders}} are a product feature. Rejecting these removes the SSTI +
    loop-bomb surface entirely, rather than relying on the Jinja sandbox alone."""


_LOGIC_MSG = "Templates may use {{placeholders}} only — not {% %} logic or {# #} comments."


def _assert_no_logic(text: str) -> None:
    if "{%" in text or "{#" in text:
        raise TemplateHasLogic(_LOGIC_MSG)


def _load_docx(data: bytes) -> DocxTemplate:
    """Parse the docx and reject Jinja statements/comments in one pass. Tags can be
    split across Word runs, so merge via docxtpl's patch_xml (the normalization it
    renders through) before scanning."""
    tpl = DocxTemplate(BytesIO(data))
    tpl.get_docx()
    _assert_no_logic(tpl.patch_xml(tpl.get_xml()))
    return tpl


def _load_pptx(data: bytes):  # type: ignore[no-untyped-def]  # pptx Presentation is Any
    """Parse the pptx and reject Jinja statements/comments in one pass."""
    prs = Presentation(BytesIO(data))
    _assert_no_logic(
        "\n".join(
            "".join(run.text for run in para.runs)
            for slide in prs.slides
            for frame in _text_frames(slide.shapes)
            for para in frame.paragraphs
        )
    )
    return prs


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------
def discover_placeholders(data: bytes, fmt: str) -> list[str]:
    """The `{{tags}}` declared in the template. docx: docxtpl's Jinja analysis
    (robust to run-splitting), sorted for determinism. pptx: scan joined
    paragraph text, de-duped in first-seen order. Both reject `{% %}`/`{# #}`
    logic (SEC-03/10) as part of the single parse."""
    assert_safe_ooxml(data)
    if fmt == "docx":
        return sorted(_load_docx(data).get_undeclared_template_variables())
    seen: list[str] = []
    for slide in _load_pptx(data).slides:
        for frame in _text_frames(slide.shapes):
            for para in frame.paragraphs:
                joined = "".join(r.text for r in para.runs)
                for name in _TAG.findall(joined):
                    if name not in seen:
                        seen.append(name)
    return seen


def _text_frames(shapes):  # type: ignore[no-untyped-def]
    """Yield every text frame in a shape tree — text boxes/placeholders + table cells."""
    for shape in shapes:
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


# ---------------------------------------------------------------------------
# Fill / render
# ---------------------------------------------------------------------------
def clamp_values(values: dict[str, str], placeholders: list[str]) -> dict[str, str]:
    """Keep only known placeholders; clamp each value's length. Missing keys are
    left out (they render blank)."""
    allowed = set(placeholders)
    return {
        k: str(v if v is not None else "")[:MAX_VALUE_LEN]
        for k, v in (values or {}).items()
        if k in allowed
    }


def render(data: bytes, values: dict[str, str], fmt: str) -> bytes:
    assert_safe_ooxml(data)
    return _fill_docx(data, values) if fmt == "docx" else _fill_pptx(data, values)


def _fill_docx(data: bytes, values: dict[str, str]) -> bytes:
    tpl = _load_docx(data)  # single parse; also rejects `{% %}`/`{# #}` logic tags
    # Sandbox the Jinja env belt-and-suspenders: _load_docx already rejects
    # statements, but keep the sandbox in case an unforeseen construct slips past.
    tpl.render(values, jinja_env=SandboxedEnvironment())  # missing var → empty string
    out = BytesIO()
    tpl.save(out)
    return out.getvalue()


def _fill_pptx(data: bytes, values: dict[str, str]) -> bytes:
    prs = _load_pptx(data)  # single parse; also rejects logic tags
    for slide in prs.slides:
        for frame in _text_frames(slide.shapes):
            for para in frame.paragraphs:
                _fill_paragraph(para, values)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def _fill_paragraph(para, values: dict[str, str]) -> None:  # type: ignore[no-untyped-def]
    """Substitute on the joined paragraph text (a tag can span runs), then write
    the result into the first run (keeping its formatting) and clear the rest."""
    runs = para.runs
    if not runs:
        return
    joined = "".join(r.text for r in runs)
    if "{{" not in joined:
        return
    replaced = _TAG.sub(lambda m: str(values.get(m.group(1), "")), joined)
    if replaced == joined:
        return
    runs[0].text = replaced
    for r in runs[1:]:
        r.text = ""


# ---------------------------------------------------------------------------
# Mongo CRUD
# ---------------------------------------------------------------------------
def _to_response(doc: dict) -> TemplateResponse:
    return TemplateResponse(**serialize_doc(doc))


async def create_template(
    db: AsyncIOMotorDatabase,
    workspace_oid: ObjectId,
    owner_oid: ObjectId,
    *,
    template_oid: ObjectId,
    name: str,
    fmt: str,
    storage_key: str,
    placeholders: list[str],
) -> TemplateResponse:
    doc = {
        "_id": template_oid,
        "workspace_id": workspace_oid,
        "owner_id": owner_oid,
        "name": name,
        "format": fmt,
        "storage_key": storage_key,
        "placeholders": placeholders,
        "created_at": utcnow(),
    }
    await db.templates.insert_one(doc)
    return _to_response(doc)


async def list_templates(
    db: AsyncIOMotorDatabase, workspace_oid: ObjectId
) -> list[TemplateResponse]:
    cursor = db.templates.find({"workspace_id": workspace_oid}).sort("created_at", -1)
    return [_to_response(doc) async for doc in cursor]


async def get_template(
    db: AsyncIOMotorDatabase, workspace_oid: ObjectId, template_id: str
) -> dict | None:
    try:
        oid = ObjectId(template_id)
    except Exception:
        return None
    return await db.templates.find_one({"_id": oid, "workspace_id": workspace_oid})


async def delete_template(db: AsyncIOMotorDatabase, doc: dict) -> None:
    await db.templates.delete_one({"_id": doc["_id"]})
