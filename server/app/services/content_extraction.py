"""Extract a bounded, lightly-structured plain-text representation of an uploaded
source document, for mapping onto a branded template's placeholders. Distinct
from `rag/loaders` (which yields Sections for chunking) — here we want one
readable blob for the fill prompt. Pure + synchronous so it's unit-testable; the
route runs it off the event loop under a timeout."""
from __future__ import annotations

import io
import os

from app.utils.archives import assert_safe_ooxml

MAX_SOURCE_CHARS = 40000  # ~10-12 pages / ~10k input tokens — generous, bounded
_MAX_PDF_PAGES = 80

_TEXT_EXTS = {".txt", ".md", ".markdown"}


class UnsupportedSourceError(ValueError):
    """The uploaded source file type can't be extracted."""


def extract_document_text(data: bytes, filename: str) -> str:
    """Return up to MAX_SOURCE_CHARS of plain text from a docx/pptx/pdf/txt/md
    source. Raises UnsupportedSourceError for other types, ValueError if no text
    could be extracted."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".docx":
        assert_safe_ooxml(data)  # zip bomb: the size cap only saw the compressed bytes
        text = _extract_docx(data)
    elif ext == ".pptx":
        assert_safe_ooxml(data)
        text = _extract_pptx(data)
    elif ext == ".pdf":
        text = _extract_pdf(data)
    elif ext in _TEXT_EXTS:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedSourceError(f"Unsupported source type: {ext or '(none)'}")
    text = text.strip()
    if not text:
        raise ValueError("No extractable text found in the document.")
    return text[:MAX_SOURCE_CHARS]


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if "heading" in style:
            lines.append(f"# {t}")
        elif "list" in style:
            lines.append(f"- {t}")
        else:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            out.append(f"Slide {i}:\n" + "\n".join(parts))
    return "\n\n".join(out)


def _extract_pdf(data: bytes) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages[:_MAX_PDF_PAGES]:
            t = page.extract_text() or ""
            if t.strip():
                pages.append(t)
    return "\n\n".join(pages)
