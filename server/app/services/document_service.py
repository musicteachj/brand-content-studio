"""Document rendering (Phase 2).

One shared structured `DocumentModel` (a dict at this layer; it round-trips
through Mongo) is rendered **independently** into each format: python-docx for
`.docx`, ReportLab for the PDF. There is no docx→PDF conversion and no headless
office/browser dependency: "matching" means both formats are built from the
same model with the same brand colors and fonts, so they agree by construction.

Brand styling: titles/headings take the workspace's primary brand color, the
eyebrow takes its accent, body text stays a readable dark. Fonts come from
`brand_profile.fonts`: the .docx references them by name; the PDF embeds the
real TTF when one is available locally (see `brand_fonts`), else falls back to a
serif/sans base font. App (Folio) fonts/colors never leak in (DESIGN §7).

The model is also clamped here (`build_document_output`) so a runaway model
response can't balloon a built file (ARCHITECTURE gotcha #9).
"""
from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from pptx.util import Pt as PptxPt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.models.generation import DocumentArchetype
from app.services.brand_fonts import font_slug, is_serif, resolve_font_file

# Bounds applied when shaping the model (defense in depth alongside the tool schema).
MAX_SECTIONS = 4
MAX_BULLETS = 6
_LIMITS = {"title": 200, "eyebrow": 80, "lede": 600, "footer": 200, "heading": 160, "body": 1200, "bullet": 200}

# Neutral document ink: the brand signal is the title/heading/accent colors +
# the brand fonts, not the body text color.
_INK = "#1A1D21"
_MUTED = "#6C737A"
_FALLBACK_PRIMARY = "#1A1D21"

# Smaller PDF type scales tried in order when a one-pager runs past one page.
_ONE_PAGER_SCALES = (0.94, 0.88, 0.82)

# Track TTFs already registered with ReportLab (registering twice raises).
_REGISTERED: set[str] = set()


# ---------------------------------------------------------------------------
# Model shaping
# ---------------------------------------------------------------------------
def _clip(value: object, limit: int) -> str:
    return str(value or "").strip()[:limit]


def build_document_output(archetype: str, extracted: dict) -> dict:
    """Shape a tool-use result into the persisted document `output`. `formats` is
    empty here: the files are built and their keys attached after the generation
    id exists (the storage key is derived from it)."""
    sections: list[dict] = []
    for raw in (extracted.get("sections") or [])[:MAX_SECTIONS]:
        if not isinstance(raw, dict):
            continue
        bullets = [
            _clip(b, _LIMITS["bullet"])
            for b in (raw.get("bullets") or [])
            if str(b).strip()
        ][:MAX_BULLETS]
        sections.append(
            {
                "heading": _clip(raw.get("heading"), _LIMITS["heading"]),
                "body": _clip(raw.get("body"), _LIMITS["body"]),
                "bullets": bullets,
            }
        )
    return {
        "document": {
            "archetype": archetype,
            "title": _clip(extracted.get("title"), _LIMITS["title"]),
            "eyebrow": _clip(extracted.get("eyebrow"), _LIMITS["eyebrow"]),
            "lede": _clip(extracted.get("lede"), _LIMITS["lede"]),
            "sections": sections,
            "footer": _clip(extracted.get("footer"), _LIMITS["footer"]),
        },
        "formats": {},
    }


# ---------------------------------------------------------------------------
# Brand color / font resolution
# ---------------------------------------------------------------------------
def _hex6(value: str | None, default: str) -> str:
    """Normalize `#rgb`/`#rrggbb` to 6 lowercase hex chars (no `#`); fall back on
    anything malformed so a bad palette value can never crash a render."""
    v = (value or "").strip().lstrip("#").lower()
    if len(v) == 3:
        v = "".join(ch * 2 for ch in v)
    if len(v) == 6 and all(c in "0123456789abcdef" for c in v):
        return v
    return default.lstrip("#").lower()


def _brand_colors(brand_profile: dict) -> dict[str, str]:
    colors = brand_profile.get("colors") or []
    by_role: dict[str, str] = {}
    first_hex: str | None = None
    for c in colors:
        if isinstance(c, dict) and isinstance(c.get("hex"), str):
            by_role.setdefault(c.get("role") or "primary", c["hex"])
            first_hex = first_hex or c["hex"]
    primary = by_role.get("primary") or first_hex or _FALLBACK_PRIMARY
    accent = (
        brand_profile.get("accent_hex")
        or by_role.get("accent")
        or by_role.get("secondary")
        or primary
    )
    return {"primary": primary, "accent": accent, "ink": _INK, "muted": _MUTED}


def _rgb(value: str, default: str) -> RGBColor:
    return RGBColor.from_string(_hex6(value, default))


def _pdf_color(value: str, default: str) -> HexColor:
    return HexColor("#" + _hex6(value, default))


def _register_pdf_font(
    family: str | None, role: str, uploaded: dict[str, bytes] | None = None
) -> str:
    """Return a ReportLab font name: an uploaded TTF (by family) or an embedded
    brand TTF when available, else a base-14 serif/sans fallback (bold for headings)."""
    src: str | BytesIO | None = None
    if uploaded and family and family in uploaded:
        src = BytesIO(uploaded[family])
    else:
        src = resolve_font_file(family)
    if src is not None:
        name = f"bcs-{role}-{font_slug(family or '')}"
        if name in _REGISTERED:
            return name
        try:
            pdfmetrics.registerFont(TTFont(name, src))
            _REGISTERED.add(name)
            return name
        except Exception:  # noqa: BLE001, S110  # a bad font file must not break rendering
            pass
    if role == "heading":
        return "Times-Bold" if is_serif(family) else "Helvetica-Bold"
    return "Times-Roman" if is_serif(family) else "Helvetica"


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------
def render_docx(document: dict, brand_profile: dict) -> bytes:
    colors = _brand_colors(brand_profile)
    fonts = brand_profile.get("fonts") or {}
    heading_font = (fonts.get("heading") or "Georgia").strip() or "Georgia"
    body_font = (fonts.get("body") or "Calibri").strip() or "Calibri"

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(11)

    def add(text: str, *, font: str, size: int, color: str, default_color: str, bold: bool = False,
            upper: bool = False, style: str | None = None) -> None:
        if not text:
            return
        para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        run = para.add_run(text.upper() if upper else text)
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color, default_color)

    add(document.get("eyebrow", ""), font=body_font, size=9, color=colors["accent"],
        default_color=_INK, bold=True, upper=True)
    add(document.get("title", ""), font=heading_font, size=26, color=colors["primary"],
        default_color=_FALLBACK_PRIMARY, bold=True)
    add(document.get("lede", ""), font=body_font, size=13, color=colors["ink"], default_color=_INK)

    for section in document.get("sections") or []:
        add(section.get("heading", ""), font=heading_font, size=15, color=colors["primary"],
            default_color=_FALLBACK_PRIMARY, bold=True)
        add(section.get("body", ""), font=body_font, size=11, color=colors["ink"], default_color=_INK)
        for bullet in section.get("bullets") or []:
            add(str(bullet), font=body_font, size=11, color=colors["ink"], default_color=_INK,
                style="List Bullet")

    add(document.get("footer", ""), font=body_font, size=9, color=colors["muted"], default_color=_MUTED)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(
    document: dict, brand_profile: dict, uploaded_fonts: dict[str, bytes] | None = None
) -> bytes:
    colors = _brand_colors(brand_profile)
    fonts = brand_profile.get("fonts") or {}
    heading_font = _register_pdf_font(fonts.get("heading"), "heading", uploaded_fonts)
    body_font = _register_pdf_font(fonts.get("body"), "body", uploaded_fonts)

    primary = _pdf_color(colors["primary"], _FALLBACK_PRIMARY)
    accent = _pdf_color(colors["accent"], _INK)
    ink = _pdf_color(colors["ink"], _INK)
    muted = _pdf_color(colors["muted"], _MUTED)

    def build(scale: float) -> tuple[bytes, int]:
        def style(name: str, font: str, size: float, color: HexColor, leading: float, **space: float) -> ParagraphStyle:
            return ParagraphStyle(
                name,
                fontName=font,
                fontSize=size * scale,
                textColor=color,
                leading=leading * scale,
                **{k: v * scale for k, v in space.items()},
            )

        eyebrow_style = style("eyebrow", body_font, 9, accent, 12, spaceAfter=4)
        title_style = style("title", heading_font, 26, primary, 30, spaceAfter=8)
        lede_style = style("lede", body_font, 13, ink, 18, spaceAfter=16)
        heading_style = style("heading", heading_font, 15, primary, 18, spaceBefore=10, spaceAfter=4)
        body_style = style("body", body_font, 11, ink, 16, spaceAfter=6)
        bullet_style = style("bullet", body_font, 11, ink, 15, leftIndent=16, spaceAfter=3)
        footer_style = style("footer", body_font, 9, muted, 12, spaceBefore=20)

        flow: list = []
        if document.get("eyebrow"):
            flow.append(Paragraph(escape(document["eyebrow"].upper()), eyebrow_style))
        if document.get("title"):
            flow.append(Paragraph(escape(document["title"]), title_style))
        if document.get("lede"):
            flow.append(Paragraph(escape(document["lede"]), lede_style))
        for section in document.get("sections") or []:
            if section.get("heading"):
                flow.append(Paragraph(escape(section["heading"]), heading_style))
            if section.get("body"):
                flow.append(Paragraph(escape(section["body"]), body_style))
            for bullet in section.get("bullets") or []:
                text = str(bullet).strip()
                if text:
                    flow.append(Paragraph("•  " + escape(text), bullet_style))
        if document.get("footer"):
            flow.append(Paragraph(escape(document["footer"]), footer_style))
        if not flow:
            flow.append(Spacer(1, 1))

        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=LETTER,
            topMargin=0.9 * inch,
            bottomMargin=0.9 * inch,
            leftMargin=1.0 * inch,
            rightMargin=1.0 * inch,
            title=document.get("title") or "Document",
        )
        doc.build(flow)
        return buf.getvalue(), doc.page

    # A one-pager is promised as a single page (QA-09): if the copy runs long, set
    # it a step smaller until it fits, down to a still-readable floor (body ~9pt).
    # If even that spills over, smaller type gains nothing: keep the full size.
    full, pages = build(1.0)
    if pages <= 1 or document.get("archetype") != DocumentArchetype.one_pager.value:
        return full
    for scale in _ONE_PAGER_SCALES:
        pdf, pages = build(scale)
        if pages <= 1:
            return pdf
    return full


def _pptx_rgb(value: str, default: str) -> PptxRGBColor:
    return PptxRGBColor.from_string(_hex6(value, default))


def render_pptx(document: dict, brand_profile: dict) -> bytes:
    colors = _brand_colors(brand_profile)
    fonts = brand_profile.get("fonts") or {}
    heading_font = (fonts.get("heading") or "Georgia").strip() or "Georgia"
    body_font = (fonts.get("body") or "Calibri").strip() or "Calibri"
    primary = _pptx_rgb(colors["primary"], _FALLBACK_PRIMARY)
    accent = _pptx_rgb(colors["accent"], _INK)
    ink = _pptx_rgb(colors["ink"], _INK)

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_line(tf, text, *, font, size, color, bold=False, upper=False, space_before=0.0):
        if not text:
            return None
        head = tf.paragraphs[0]
        # Reuse the text frame's default (empty) first paragraph for the first line.
        para = head if len(tf.paragraphs) == 1 and not head.runs else tf.add_paragraph()
        run = para.add_run()
        run.text = text.upper() if upper else text
        run.font.name = font
        run.font.size = PptxPt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        if space_before:
            para.space_before = PptxPt(space_before)
        return para

    def new_box(top: float, height: float):
        box = prs.slides.add_slide(blank).shapes.add_textbox(
            Inches(0.9), Inches(top), Inches(11.5), Inches(height)
        )
        box.text_frame.word_wrap = True
        return box.text_frame

    # Title slide: title + eyebrow + lede.
    tf = new_box(2.2, 3.1)
    add_line(tf, document.get("eyebrow"), font=body_font, size=16, color=accent, bold=True, upper=True)
    add_line(tf, document.get("title") or "Untitled", font=heading_font, size=40, color=primary, bold=True, space_before=6)
    add_line(tf, document.get("lede"), font=body_font, size=18, color=ink, space_before=12)

    # One content slide per section.
    for section in document.get("sections") or []:
        tf = new_box(0.8, 5.9)
        add_line(tf, section.get("heading"), font=heading_font, size=28, color=primary, bold=True)
        add_line(tf, section.get("body"), font=body_font, size=18, color=ink, space_before=10)
        for bullet in section.get("bullets") or []:
            text = str(bullet).strip()
            if text:
                add_line(tf, f"•  {text}", font=body_font, size=18, color=ink, space_before=6)

    # Closing slide: the footer line, centered.
    if document.get("footer"):
        tf = new_box(3.2, 1.3)
        para = add_line(tf, document["footer"], font=heading_font, size=24, color=primary, bold=True)
        if para is not None:
            para.alignment = PP_ALIGN.CENTER

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_document_files(
    document: dict, brand_profile: dict, uploaded_fonts: dict[str, bytes] | None = None
) -> dict[str, bytes]:
    """Render every format an archetype ships from the one shared model:
    one-pager → `.docx` + PDF; short-deck → `.pptx`."""
    if document.get("archetype") == DocumentArchetype.short_deck.value:
        return {"pptx": render_pptx(document, brand_profile)}
    return {
        "docx": render_docx(document, brand_profile),
        "pdf": render_pdf(document, brand_profile, uploaded_fonts),
    }
