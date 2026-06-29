"""Extract bounded plain text from an uploaded source document (docx/pptx/pdf/
text). Office fixtures are built in-memory; no model calls."""
import io

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.services.content_extraction import (
    MAX_SOURCE_CHARS,
    UnsupportedSourceError,
    extract_document_text,
)


def _docx(*, paras=(), heading=None) -> bytes:
    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for p in paras:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx(*lines: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.add_run().text = line
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_extract_docx_keeps_heading_and_body():
    out = extract_document_text(
        _docx(heading="The Plan", paras=["First point.", "Second point."]), "plan.docx"
    )
    assert "# The Plan" in out
    assert "First point." in out and "Second point." in out


def test_extract_pptx_prefixes_slides():
    out = extract_document_text(_pptx("Title slide", "a bullet"), "deck.pptx")
    assert "Slide 1:" in out
    assert "Title slide" in out and "a bullet" in out


def test_extract_text_and_markdown():
    assert "hello world" in extract_document_text(b"hello world", "notes.txt")
    assert "# Heading" in extract_document_text(b"# Heading\n\nbody", "notes.md")


def test_extract_truncates_to_budget():
    big = ("word " * 20000).encode("utf-8")  # ~100k chars
    out = extract_document_text(big, "big.txt")
    assert len(out) == MAX_SOURCE_CHARS


def test_extract_rejects_unsupported_type():
    with pytest.raises(UnsupportedSourceError):
        extract_document_text(b"\x00\x01", "image.png")


def test_extract_empty_raises_valueerror():
    with pytest.raises(ValueError):
        extract_document_text(_docx(paras=["   "]), "blank.docx")
