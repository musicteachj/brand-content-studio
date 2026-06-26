"""Phase 2 — document output: a branded .docx + matching PDF from one model.

The python-docx style assertions are the backbone (PLAN.md Phase 2): a built
.docx carries the workspace's heading/body fonts and primary brand color, never
the app's Folio chrome. Plus the SSE route (persists + builds files) and the
auth-gated, ownership-scoped download route.
"""
import io
import json

import pytest
import pytest_asyncio
from bson import ObjectId
from docx import Document
from docx.shared import RGBColor
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor

import app.routes.documents as documents_mod
from app.config import settings
from app.database import get_database
from app.main import app
from app.services import document_service, storage_service
from app.services.generation_service import GenerationService
from app.services.storage_service import build_document_key, get_storage
from app.utils.dependencies import get_workspace_or_404
from tests.conftest import NS_A, TEST_DB_NAME, build_corpus_chunks
from tests.fakes import FakeGenerationService, FakeVectorStore

WORKSPACE_OID = ObjectId()
OWNER_OID = ObjectId()

BRAND_PROFILE = {
    "colors": [
        {"name": "Forest", "hex": "#1C3A2E", "role": "primary"},
        {"name": "Cream", "hex": "#F4F1E9", "role": "neutral"},
        {"name": "Moss", "hex": "#6B8F71", "role": "accent"},
    ],
    "fonts": {"heading": "DM Serif Display", "body": "Spectral"},
    "voice_summary": "Grounded, data-first voice.",
}

FIXED_MODEL = {
    "archetype": "one-pager",
    "title": "Recycled-Denim Insulation",
    "eyebrow": "Product brief",
    "lede": "A high-performance insulation made from post-consumer denim.",
    "sections": [
        {
            "heading": "Why it matters",
            "body": "It cuts embodied carbon across the wall assembly.",
            "bullets": ["Lower embodied carbon", "Safe to handle"],
        },
    ],
    "footer": "Verdure Materials",
}


# --- renderer unit tests (no DB / no network) ----------------------------
def test_render_docx_applies_brand_styles():
    data = document_service.render_docx(FIXED_MODEL, BRAND_PROFILE)
    doc = Document(io.BytesIO(data))
    runs = [r for p in doc.paragraphs for r in p.runs]
    texts = [p.text for p in doc.paragraphs]

    # Title uses the brand heading font + primary brand color.
    title = next(r for r in runs if r.text == "Recycled-Denim Insulation")
    assert title.font.name == "DM Serif Display"
    assert title.font.color.rgb == RGBColor.from_string("1C3A2E")

    # Body uses the brand body font (never an app font).
    body = next(r for r in runs if "embodied carbon" in r.text)
    assert body.font.name == "Spectral"

    # Structure present: section heading + bullets.
    assert "Why it matters" in texts
    assert any("Lower embodied carbon" in t for t in texts)

    # No Folio app fonts leaked in (DESIGN §7).
    used_fonts = {r.font.name for r in runs if r.font.name}
    assert "Plus Jakarta Sans" not in used_fonts
    assert "DM Mono" not in used_fonts


def test_render_pdf_is_valid_pdf():
    data = document_service.render_pdf(FIXED_MODEL, BRAND_PROFILE)
    assert data[:5] == b"%PDF-"
    assert len(data) > 800  # a non-trivial, built document


def test_render_pptx_is_a_branded_deck():
    model = {**FIXED_MODEL, "archetype": "short-deck"}
    data = document_service.render_pptx(model, BRAND_PROFILE)
    assert data[:2] == b"PK"  # .pptx is a zip container

    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    # Title slide + one per section + a closing footer slide.
    assert len(slides) == len(model["sections"]) + 2

    runs = [r for s in slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    title = next(r for r in runs if r.text == "Recycled-Denim Insulation")
    assert title.font.name == "DM Serif Display"
    assert title.font.color.rgb == PptxRGBColor.from_string("1C3A2E")
    fonts = {r.font.name for r in runs if r.font.name}
    assert "Plus Jakarta Sans" not in fonts and "DM Mono" not in fonts


def test_build_document_output_clamps_runaway_model():
    extracted = {
        "title": "T",
        "eyebrow": "E",
        "lede": "L",
        "footer": "F",
        "sections": [
            {"heading": f"H{i}", "body": "B", "bullets": [f"b{j}" for j in range(20)]}
            for i in range(8)
        ],
    }
    out = document_service.build_document_output("one-pager", extracted)
    assert len(out["document"]["sections"]) <= document_service.MAX_SECTIONS
    assert all(
        len(s["bullets"]) <= document_service.MAX_BULLETS
        for s in out["document"]["sections"]
    )
    assert out["formats"] == {}  # keys attached only after the files are built


# --- route + download tests ----------------------------------------------
def _workspace(ready_chunk_count: int) -> dict:
    return {
        "_id": WORKSPACE_OID,
        "owner_id": OWNER_OID,
        "vector_namespace": NS_A,
        "ready_chunk_count": ready_chunk_count,
        "brand_profile": BRAND_PROFILE,
    }


@pytest_asyncio.fixture
async def client(test_db, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "database_name", TEST_DB_NAME)
    monkeypatch.setattr(settings, "storage_local_dir", str(tmp_path))
    storage_service.get_storage.cache_clear()

    store = FakeVectorStore()
    await store.upsert_chunks(NS_A, build_corpus_chunks(NS_A))
    monkeypatch.setattr(documents_mod, "get_vector_store", lambda: store)

    fake_gen = FakeGenerationService()
    app.dependency_overrides[get_database] = lambda: test_db
    app.dependency_overrides[GenerationService] = lambda: fake_gen
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        ac.fake_gen = fake_gen  # type: ignore[attr-defined]
        yield ac
    app.dependency_overrides.clear()
    storage_service.get_storage.cache_clear()


def _override_workspace(ready_chunk_count: int) -> None:
    app.dependency_overrides[get_workspace_or_404] = lambda: _workspace(ready_chunk_count)


DOC_URL = f"/api/workspaces/{WORKSPACE_OID}/generate/document"
GEN_BASE = f"/api/workspaces/{WORKSPACE_OID}/generations"
_BRIEF = {"brief": "A one-pager on recycled denim insulation", "archetype": "one-pager", "tone": "default"}


def _meta(resp_text: str) -> dict:
    payloads = [
        json.loads(line[len("data: ") :])
        for line in resp_text.splitlines()
        if line.startswith("data: ") and line[len("data: ") :] != "[DONE]"
    ]
    metas = [p["meta"] for p in payloads if "meta" in p]
    assert len(metas) == 1, "exactly one terminal meta frame"
    return metas[0]


@pytest.mark.asyncio
async def test_empty_corpus_blocks_document(client):
    _override_workspace(0)
    resp = await client.post(DOC_URL, json=_BRIEF)
    assert resp.status_code == 409
    assert client.fake_gen.calls == []  # Claude must not have been called


@pytest.mark.asyncio
async def test_generate_document_persists_and_builds_files(client, test_db):
    _override_workspace(12)
    resp = await client.post(DOC_URL, json=_BRIEF)
    assert resp.status_code == 200

    assert "[DONE]" in resp.text
    assert '"text"' in resp.text, "draft streamed as text frames"
    meta = _meta(resp.text)
    assert meta["archetype"] == "one-pager"
    assert meta["document"]["title"]
    assert set(meta["formats"]) == {"docx", "pdf"}
    assert client.fake_gen.doc_section_calls, "tool-use extraction was called"

    doc = await test_db.generations.find_one({"_id": ObjectId(meta["generation_id"])})
    assert doc is not None
    assert doc["type"] == "document"
    assert doc["status"] == "complete"
    assert set(doc["output"]["formats"]) == {"docx", "pdf"}
    assert doc["retrieved_chunk_ids"], "retrieval snapshot persisted"

    # Files exist in storage and are the right kind.
    gid = meta["generation_id"]
    docx_bytes = await get_storage().read(build_document_key(str(WORKSPACE_OID), gid, "docx"))
    pdf_bytes = await get_storage().read(build_document_key(str(WORKSPACE_OID), gid, "pdf"))
    assert docx_bytes[:2] == b"PK"  # .docx is a zip container
    assert pdf_bytes[:5] == b"%PDF-"


@pytest.mark.asyncio
async def test_download_route_serves_and_scopes(client):
    _override_workspace(12)
    resp = await client.post(DOC_URL, json=_BRIEF)
    gid = _meta(resp.text)["generation_id"]

    docx = await client.get(f"{GEN_BASE}/{gid}/download/docx")
    assert docx.status_code == 200
    assert docx.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in docx.headers["content-disposition"]
    assert docx.content[:2] == b"PK"

    pdf = await client.get(f"{GEN_BASE}/{gid}/download/pdf")
    assert pdf.status_code == 200
    assert pdf.content[:5] == b"%PDF-"

    # Unknown format and unknown generation both 404 (no leak, no traversal).
    assert (await client.get(f"{GEN_BASE}/{gid}/download/exe")).status_code == 404
    assert (await client.get(f"{GEN_BASE}/{ObjectId()}/download/pdf")).status_code == 404


@pytest.mark.asyncio
async def test_short_deck_builds_and_serves_pptx(client, test_db):
    _override_workspace(12)
    resp = await client.post(
        DOC_URL,
        json={"brief": "A short deck pitching recycled denim insulation", "archetype": "short-deck", "tone": "default"},
    )
    assert resp.status_code == 200
    meta = _meta(resp.text)
    assert meta["archetype"] == "short-deck"
    # A deck ships only .pptx — not the one-pager's docx/pdf.
    assert set(meta["formats"]) == {"pptx"}

    doc = await test_db.generations.find_one({"_id": ObjectId(meta["generation_id"])})
    assert doc["type"] == "document"
    assert set(doc["output"]["formats"]) == {"pptx"}

    gid = meta["generation_id"]
    pptx = await client.get(f"{GEN_BASE}/{gid}/download/pptx")
    assert pptx.status_code == 200
    assert pptx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml"
    )
    assert pptx.content[:2] == b"PK"
    # The formats this archetype didn't build aren't downloadable.
    assert (await client.get(f"{GEN_BASE}/{gid}/download/docx")).status_code == 404
