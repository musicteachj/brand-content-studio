"""On-brand document templates (Phase 5). Discovery + fill are pure (fixtures
built in-memory with python-docx/python-pptx); routes are tested with the AI
fill faked (no model call)."""
import io

import pytest
import pytest_asyncio
from bson import ObjectId
from docx import Document
from docx.shared import Pt, RGBColor
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.util import Inches, Pt as PptxPt

from app.config import settings
from app.database import get_database
from app.main import app
from app.models.common import utcnow
from app.models.template import TemplateAiFillRequest, TemplateRenderRequest, TemplateResponse
from app.models.user import UserInDB
from app.prompts.template_fill import (
    build_fill_system,
    build_fill_tool,
    build_fill_user,
    build_fill_user_from_content,
)
from app.services.generation_service import GenerationService
from app.services.storage_service import build_template_key, get_storage
from app.services.template_service import (
    clamp_values,
    discover_placeholders,
    render,
)
from app.utils.dependencies import get_current_user
from tests.conftest import TEST_DB_NAME


# --- Task 1: storage key ----------------------------------------------------

def test_build_template_key_is_server_derived():
    key = build_template_key("ws1", "t1", "Quarterly Report.docx")
    assert key == "ws1/templates/t1/Quarterly_Report.docx"


# --- Task 2: discovery ------------------------------------------------------

def _docx_bytes(*paragraphs: str) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(*lines: str) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf = box.text_frame
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.add_run().text = line
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_discover_placeholders_docx():
    data = _docx_bytes("Hello {{client_name}}", "Re: {{topic}} and {{topic}}")
    assert set(discover_placeholders(data, "docx")) == {"client_name", "topic"}


def test_discover_placeholders_pptx_in_order():
    data = _pptx_bytes("{{title}}", "by {{author}}")
    assert discover_placeholders(data, "pptx") == ["title", "author"]


# --- Task 3: fill / render / clamp ------------------------------------------

def _styled_docx() -> bytes:
    doc = Document()
    run = doc.add_paragraph().add_run("{{headline}}")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1C, 0x3A, 0x2E)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_render_docx_swaps_text_preserves_style():
    out = render(_styled_docx(), {"headline": "Carbon in the wall"}, "docx")
    doc = Document(io.BytesIO(out))
    run = doc.paragraphs[0].runs[0]
    assert run.text == "Carbon in the wall"
    assert run.font.size == Pt(22)
    assert run.font.color.rgb == RGBColor(0x1C, 0x3A, 0x2E)


def test_render_pptx_swaps_fragmented_tag_preserves_style():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    para = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.paragraphs[0]
    for frag in ("{{", "title", "}}"):
        r = para.add_run()
        r.text = frag
        r.font.size = PptxPt(40)
    buf = io.BytesIO()
    prs.save(buf)

    out = render(buf.getvalue(), {"title": "Verdure"}, "pptx")
    p2 = Presentation(io.BytesIO(out)).slides[0].shapes[0].text_frame.paragraphs[0]
    assert "".join(r.text for r in p2.runs) == "Verdure"
    assert p2.runs[0].font.size == PptxPt(40)


def test_clamp_values_drops_unknown_and_clamps_length():
    out = clamp_values({"a": "x" * 5000, "z": "nope"}, ["a"])
    assert set(out) == {"a"}
    assert len(out["a"]) == 2000


# --- Task 4: models ---------------------------------------------------------

def test_template_models_shape():
    t = TemplateResponse(id="t1", workspace_id="w1", name="Report", format="docx",
                         storage_key="k", placeholders=["title"])
    assert t.format == "docx" and t.placeholders == ["title"]
    assert TemplateRenderRequest(values={"title": "x"}).values == {"title": "x"}
    assert TemplateAiFillRequest(brief="a deck about hempcrete").brief


# --- Task 6: AI-fill prompt/tool --------------------------------------------

def test_build_fill_tool_uses_placeholder_names():
    tool = build_fill_tool(["title", "body"])
    props = tool["input_schema"]["properties"]
    assert set(props) == {"title", "body"}
    assert tool["input_schema"]["required"] == ["title", "body"]
    assert tool["strict"] is True


def test_fill_prompt_builders_nonempty():
    assert build_fill_system().strip()
    assert "title" in build_fill_user("a brief", ["title"])


def test_build_fill_user_from_content_includes_content_and_fields():
    out = build_fill_user_from_content("The Q3 plan: grow EU sales.", ["title", "summary"])
    assert "The Q3 plan: grow EU sales." in out
    assert "title" in out and "summary" in out


# --- Task 7: routes (AI fill faked) -----------------------------------------

WS_OID = ObjectId()
OWNER_OID = ObjectId()
DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class _FakeGen:
    async def fill_template_values(self, system, user, tool):
        return {name: f"draft {name}" for name in tool["input_schema"]["properties"]}


def _user(is_demo: bool = False) -> UserInDB:
    return UserInDB(id=str(OWNER_OID), email="o@example.com", hashed_password="x",
                    full_name="O", is_demo=is_demo)


async def _seed_ws(test_db):
    await test_db.workspaces.insert_one({
        "_id": WS_OID, "owner_id": OWNER_OID, "name": "Verdure", "slug": "verdure",
        "brand_profile": {}, "vector_namespace": str(WS_OID),
        "embedding_model": settings.embedding_model, "embedding_dim": settings.embedding_dim,
        "asset_count": 0, "ready_chunk_count": 0, "created_at": utcnow(), "updated_at": utcnow(),
    })


async def _mk(test_db, tmp_path, monkeypatch, *, is_demo=False) -> AsyncClient:
    monkeypatch.setattr(settings, "database_name", TEST_DB_NAME)
    monkeypatch.setattr(settings, "storage_local_dir", str(tmp_path))
    get_storage.cache_clear()
    await _seed_ws(test_db)
    app.dependency_overrides[get_database] = lambda: test_db
    app.dependency_overrides[get_current_user] = lambda: _user(is_demo)
    app.dependency_overrides[GenerationService] = lambda: _FakeGen()
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


@pytest_asyncio.fixture
async def tpl_client(test_db, tmp_path, monkeypatch):
    ac = await _mk(test_db, tmp_path, monkeypatch)
    async with ac:
        yield ac
    app.dependency_overrides.clear()
    get_storage.cache_clear()


def _base() -> str:
    return f"/api/workspaces/{WS_OID}/templates"


async def _upload(ac) -> dict:
    data = _docx_bytes("Hello {{client_name}}")
    return (await ac.post(_base(), files={"file": ("Report.docx", data, DOCX_CT)})).json()


@pytest.mark.asyncio
async def test_upload_discovers_and_stores(tpl_client, test_db):
    body = await _upload(tpl_client)
    assert body["format"] == "docx"
    assert body["placeholders"] == ["client_name"]
    assert await get_storage().read(body["storage_key"])
    assert await test_db.templates.find_one({"_id": ObjectId(body["id"])})


@pytest.mark.asyncio
async def test_upload_rejects_bad_type(tpl_client):
    resp = await tpl_client.post(_base(), files={"file": ("notes.txt", b"hi", "text/plain")})
    assert resp.status_code == 415


@pytest.mark.asyncio
async def test_list_and_delete(tpl_client, test_db):
    t = await _upload(tpl_client)
    listed = (await tpl_client.get(_base())).json()
    assert len(listed) == 1
    resp = await tpl_client.delete(f"{_base()}/{t['id']}")
    assert resp.status_code == 200
    assert await test_db.templates.find_one({"_id": ObjectId(t["id"])}) is None


@pytest.mark.asyncio
async def test_ai_fill_returns_values(tpl_client):
    t = await _upload(tpl_client)
    resp = await tpl_client.post(f"{_base()}/{t['id']}/ai-fill", json={"brief": "a welcome letter"})
    assert resp.status_code == 200
    assert resp.json()["values"] == {"client_name": "draft client_name"}


@pytest.mark.asyncio
async def test_fill_from_document_maps_content(tpl_client):
    t = await _upload(tpl_client)
    src = _docx_bytes("Welcome to Acme", "We supply hempcrete.")
    resp = await tpl_client.post(
        f"{_base()}/{t['id']}/fill-from-document",
        files={"file": ("draft.docx", src, DOCX_CT)},
    )
    assert resp.status_code == 200
    assert resp.json()["values"] == {"client_name": "draft client_name"}


@pytest.mark.asyncio
async def test_fill_from_document_rejects_bad_type(tpl_client):
    t = await _upload(tpl_client)
    resp = await tpl_client.post(
        f"{_base()}/{t['id']}/fill-from-document",
        files={"file": ("logo.png", b"\x89PNG", "image/png")},
    )
    assert resp.status_code == 415


@pytest.mark.asyncio
async def test_fill_from_document_missing_template_404(tpl_client):
    resp = await tpl_client.post(
        f"{_base()}/{ObjectId()}/fill-from-document",
        files={"file": ("draft.docx", _docx_bytes("hi"), DOCX_CT)},
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_fill_from_document_allowed_for_demo(tpl_client):
    t = await _upload(tpl_client)
    # Re-point auth at a demo user (same owner id) — the route is demo-allowed.
    app.dependency_overrides[get_current_user] = lambda: _user(is_demo=True)
    resp = await tpl_client.post(
        f"{_base()}/{t['id']}/fill-from-document",
        files={"file": ("draft.docx", _docx_bytes("hi"), DOCX_CT)},
    )
    assert resp.status_code == 200


@pytest.mark.asyncio
async def test_render_returns_filled_file(tpl_client):
    t = await _upload(tpl_client)
    resp = await tpl_client.post(f"{_base()}/{t['id']}/render", json={"values": {"client_name": "Acme"}})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == DOCX_CT
    text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
    assert "Hello Acme" in text


@pytest.mark.asyncio
async def test_upload_demo_blocked(test_db, tmp_path, monkeypatch):
    ac = await _mk(test_db, tmp_path, monkeypatch, is_demo=True)
    async with ac:
        resp = await ac.post(_base(), files={"file": ("Report.docx", _docx_bytes("{{x}}"), DOCX_CT)})
    app.dependency_overrides.clear()
    get_storage.cache_clear()
    assert resp.status_code == 403
